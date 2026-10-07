import os
import sys
sys.path.insert(0, os.path.abspath("."))
import pandas as pd
import openpyxl
from utils.research import demo_dataset
from utils.validation import clean_dataframe, add_validation_scores, add_lead_score
from utils.exporter import to_excel_bytes

print("1. Loading demo dataset...")
raw = demo_dataset()
print(f"Raw shape: {raw.shape}")
assert len(raw) == 25, f"Expected 25 demo records, got {len(raw)}"

print("2. Cleaning dataframe...")
clean, stats = clean_dataframe(raw)
print(f"Cleaned shape: {clean.shape}, stats: {stats}")
assert stats['duplicates_detected'] == 3, f"Expected 3 duplicates detected, got {stats['duplicates_detected']}"

print("3. Scoring dataframe...")
scored = add_validation_scores(clean)
scored = add_lead_score(scored, target_industry="Artificial Intelligence", target_location="Mumbai")

summary = {
    'Query': 'Artificial Intelligence companies in Mumbai',
    'Target Industry': 'Artificial Intelligence',
    'Target Location': 'Mumbai',
    'Target Company Type': 'B2B SaaS / Product',
    'Keywords': 'automation',
    'Generated At': '2026-10-07 22:30:00',
    'Total Raw Records Ingested': stats.get('total_raw_records', len(raw)),
    'Valid Unique Companies': len(scored),
    'High Potential Leads': int((scored['Lead Category'] == 'High Potential').sum()),
    'Medium Potential Leads': int((scored['Lead Category'] == 'Medium Potential').sum()),
    'Low Potential Leads': int((scored['Lead Category'] == 'Low Potential').sum()),
    'Average Lead Score': round(scored['Lead Score'].mean(), 1),
    'Average Confidence Score': round(scored['Confidence Score'].mean(), 1),
    'Average Data Completeness %': round(scored['Data Completeness %'].mean(), 1),
}

print("4. Generating Excel bytes...")
excel_bytes = to_excel_bytes(scored, summary=summary, raw_df=raw, stats=stats)
print(f"Excel bytes generated: {len(excel_bytes)} bytes")

# Save to disk and re-open with openpyxl to verify
excel_test_path = "scratch/test_report.xlsx"
os.makedirs("scratch", exist_ok=True)
with open(excel_test_path, "wb") as f:
    f.write(excel_bytes)

wb = openpyxl.load_workbook(excel_test_path)
print(f"Sheet names: {wb.sheetnames}")
expected_sheets = ["Research Results", "Data Quality", "Lead Scoring", "Research Summary"]
assert wb.sheetnames == expected_sheets, f"Expected {expected_sheets}, got {wb.sheetnames}"

ws1 = wb["Research Results"]
print(f"Sheet 1 rows: {ws1.max_row}, cols: {ws1.max_column}, freeze_panes: {ws1.freeze_panes}")
assert ws1.freeze_panes == "C2", f"Expected C2, got {ws1.freeze_panes}"
assert ws1.auto_filter.ref is not None, "Auto filter missing in Sheet 1"

ws2 = wb["Data Quality"]
print(f"Sheet 2 rows: {ws2.max_row}, cols: {ws2.max_column}")
assert ws2.cell(row=1, column=1).value is not None

ws3 = wb["Lead Scoring"]
print(f"Sheet 3 rows: {ws3.max_row}, cols: {ws3.max_column}, freeze_panes: {ws3.freeze_panes}")
assert ws3.freeze_panes == "D14", f"Expected D14, got {ws3.freeze_panes}"

ws4 = wb["Research Summary"]
print(f"Sheet 4 rows: {ws4.max_row}, cols: {ws4.max_column}")
assert ws4.cell(row=1, column=1).value is not None

print("ALL VERIFICATIONS PASSED SUCCESSFULLY!")
